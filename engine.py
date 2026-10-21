import io
import json
import re
import unicodedata
import pandas as pd


def normalize(value):
    if pd.isna(value):
        return ''
    value = ''.join(c for c in unicodedata.normalize('NFKD', str(value)) if not unicodedata.combining(c)).upper()
    return re.sub(r'\s+', ' ', re.sub(r'[^A-Z0-9&\s]', ' ', value)).strip()


def validate_profile(profile):
    if not isinstance(profile, dict):
        raise ValueError('El perfil debe ser un objeto JSON.')
    for key in ('perfil_id', 'codigo_curso', 'nombre_curso'):
        if not isinstance(profile.get(key), str) or not profile[key].strip():
            raise ValueError(f'Falta {key}.')
    for dimension, key in [('Cargo', 'keywords_cargo'), ('Empresa', 'keywords_empresa'), ('Profesion', 'keywords_profesion')]:
        for priority in ('P1', 'P2', 'P3'):
            words = profile.get(key, {}).get(priority)
            if not isinstance(words, list) or any(not isinstance(w, str) or not normalize(w) for w in words):
                raise ValueError(f'{key}.{priority} debe ser una lista de palabras.')
            points = profile.get('PUNTOS', {}).get(dimension, {}).get(priority)
            if type(points) not in (int, float) or not 0 <= points <= 100:
                raise ValueError(f'Puntaje inválido: {dimension}.{priority}.')
    for key in ('keywords_nivel',):
        if not isinstance(profile.get(key), list) or any(not isinstance(w, str) or not normalize(w) for w in profile[key]):
            raise ValueError(f'{key} debe ser una lista de palabras.')
    exclusions = profile.get('keywords_excluir', {})
    if not isinstance(exclusions, dict) or any(not isinstance(v, list) or any(not isinstance(w, str) or not normalize(w) for w in v) for v in exclusions.values()):
        raise ValueError('Exclusiones inválidas.')
    if type(profile.get('BONO_NIVEL')) not in (int, float) or not 0 <= profile['BONO_NIVEL'] <= 100:
        raise ValueError('Bono de nivel inválido.')
    if profile.get('UMBRALES') != [[70, 'A'], [45, 'B'], [20, 'C'], [0, 'D']]:
        raise ValueError('Esta versión requiere umbrales A≥70, B≥45, C≥20 y D≥0.')
    return profile


def pattern(words):
    parts = [r'\b' + re.escape(normalize(w)) + (r'\b' if len(normalize(w)) <= 3 else '') for w in words]
    return re.compile('|'.join(parts) if parts else r'(?!)')


def read_csv(data):
    for encoding in ('utf-8-sig', 'utf-8', 'latin-1'):
        try:
            return pd.read_csv(io.BytesIO(data), encoding=encoding, sep=None, engine='python', dtype=str)
        except UnicodeDecodeError:
            continue
    raise ValueError('No se pudo leer el CSV.')


def process(df, profile, columns, apply_exclusions=True):
    validate_profile(profile)
    regex = {dim: {p: pattern(profile[key][p]) for p in ('P1', 'P2', 'P3')} for dim, key in [('Cargo', 'keywords_cargo'), ('Empresa', 'keywords_empresa'), ('Profesion', 'keywords_profesion')]}
    level = pattern(profile['keywords_nivel'])
    exclusions = {reason: pattern(words) for reason, words in profile.get('keywords_excluir', {}).items()}
    records = []
    for _, row in df.iterrows():
        texts = {dim: normalize(row[columns[dim]]) for dim in regex}
        result = {}
        score = 0
        for dim in regex:
            priority, matches, points = '-', [], 0
            for p, rx in regex[dim].items():
                matches = rx.findall(texts[dim])
                if matches:
                    priority, points = p, profile['PUNTOS'][dim][p]
                    break
            result[f'Prioridad_{dim}'] = priority
            result[f'Match_{dim}'] = ', '.join(sorted(set(matches)))
            score += points
        bonus = profile['BONO_NIVEL'] if level.search(texts['Cargo']) else 0
        reason = ''
        if apply_exclusions:
            for name, rx in exclusions.items():
                text = texts['Cargo'] + ' ' + texts['Profesion'] if name.startswith('SST') else texts['Cargo']
                if rx.search(text):
                    reason = name
                    break
        result.update(Score_Prioridad=score + bonus, Bono_Nivel=bonus, Motivo_Exclusion=reason,
                      Segmento='EXCLUIDO' if reason else next(seg for minimum, seg in profile['UMBRALES'] if score + bonus >= minimum))
        records.append(result)
    generated = list(records[0]) if records else []
    if set(generated) & set(df.columns):
        raise ValueError('El CSV ya contiene columnas de resultados. Carga la base original.')
    out = pd.concat([df.reset_index(drop=True), pd.DataFrame(records)], axis=1)
    before = len(out)
    email = columns.get('Email')
    if email:
        out['_email'] = out[email].fillna('').str.strip().str.lower()
        out = out.sort_values('Score_Prioridad', ascending=False, kind='stable')
        out = pd.concat([out[out['_email'] != ''].drop_duplicates('_email'), out[out['_email'] == '']]).drop(columns='_email')
    out = out.sort_values(['Segmento', 'Score_Prioridad', 'Prioridad_Cargo', 'Prioridad_Empresa'],
        key=lambda s: s.map({'A':0,'B':1,'C':2,'D':3,'EXCLUIDO':4}) if s.name == 'Segmento' else -s if s.name == 'Score_Prioridad' else s.map({'P1':0,'P2':1,'P3':2,'-':3}), kind='stable')
    return out.reset_index(drop=True), before - len(out)


def excel_bytes(df, company):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    stream = io.BytesIO()
    prioritized = df[df.Segmento.isin(['A','B','C'])]
    excluded = df[df.Segmento == 'EXCLUIDO']
    summary = df.Segmento.value_counts().reindex(['A','B','C','D','EXCLUIDO'], fill_value=0).rename_axis('Segmento').reset_index(name='Registros')
    summary['% del total'] = (summary.Registros / max(len(df), 1) * 100).round(1)
    with pd.ExcelWriter(stream, engine='openpyxl') as writer:
        prioritized.to_excel(writer, sheet_name='Base Priorizada', index=False)
        summary.to_excel(writer, sheet_name='Resumen', index=False)
        prioritized[company].fillna('(vacío)').value_counts().head(30).rename_axis('Empresa').reset_index(name='Prospectos A-C').to_excel(writer, sheet_name='Resumen', startcol=4, index=False)
        excluded.Motivo_Exclusion.value_counts().rename_axis('Motivo_Exclusion').reset_index(name='Registros').to_excel(writer, sheet_name='Resumen', startcol=8, index=False)
        excluded.to_excel(writer, sheet_name='Excluidos', index=False)
        df.to_excel(writer, sheet_name='Base Completa con Score', index=False)
        for ws in writer.sheets.values():
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='1F3864')
            for i, col in enumerate(ws.columns, 1):
                ws.column_dimensions[get_column_letter(i)].width = min(45, max(12, max((len(str(c.value or '')) for c in list(col)[:500]), default=10) + 2))
            headers = [c.value for c in ws[1]]
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if cell.data_type == 'f':
                        cell.data_type = 's'
                if 'Segmento' in headers:
                    cell = row[headers.index('Segmento')]
                    cell.fill = PatternFill('solid', fgColor={'A':'C6EFCE','B':'FFEB9C','C':'FCE4D6','D':'EDEDED','EXCLUIDO':'F4CCCC'}.get(cell.value,'FFFFFF'))
    return stream.getvalue()
